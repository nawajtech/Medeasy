"""Generate ApnaMedi product overview PowerPoint."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_DARK = RGBColor(0x0F, 0x76, 0x6E)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
MINT = RGBColor(0x99, 0xF6, 0xE4)
MINT_SOFT = RGBColor(0xCC, 0xFB, 0xF1)
NAVY = RGBColor(0x0E, 0x74, 0x90)

TOTAL = 12


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def para(
    tf,
    text,
    size=16,
    bold=False,
    color=INK,
    space_before=0,
    space_after=6,
    align=PP_ALIGN.LEFT,
    font="Calibri",
):
    if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size, bold, color, font)
    return p


def footer(slide, page, total=TOTAL):
    add_rect(slide, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35), TEAL_DARK)
    _, tf = textbox(slide, Inches(0.4), Inches(7.18), Inches(10), Inches(0.3))
    para(tf, "ApnaMedi  |  Confidential", size=11, color=WHITE, space_after=0)
    _, tf2 = textbox(slide, Inches(11.2), Inches(7.18), Inches(1.8), Inches(0.3))
    para(tf2, f"{page} / {total}", size=11, color=WHITE, align=PP_ALIGN.RIGHT, space_after=0)


def section_header(slide, title, subtitle=None):
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.15), TEAL_DARK)
    _, tf = textbox(slide, Inches(0.5), Inches(0.25), Inches(12), Inches(0.5))
    para(tf, title, size=28, bold=True, color=WHITE, space_after=0)
    if subtitle:
        _, tf2 = textbox(slide, Inches(0.5), Inches(0.7), Inches(12), Inches(0.35))
        para(tf2, subtitle, size=14, color=MINT_SOFT, space_after=0)


# ===== 1 Title =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), TEAL_DARK)
add_rect(s, Inches(0), Inches(0), Inches(0.25), Inches(7.5), TEAL)
_, tf = textbox(s, Inches(0.8), Inches(2.0), Inches(11.5), Inches(1))
para(tf, "ApnaMedi", size=54, bold=True, color=WHITE, space_after=0)
_, tf = textbox(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(0.6))
para(tf, "Healthcare Management SaaS Platform", size=26, color=MINT, space_after=0)
_, tf = textbox(s, Inches(0.8), Inches(4.0), Inches(11.5), Inches(0.8))
para(tf, "Product Overview Presentation", size=18, color=WHITE, space_after=8)
para(
    tf,
    "Clinics  ·  Hospitals  ·  Pathology Labs  ·  Diagnostic Centers",
    size=15,
    color=MINT_SOFT,
    space_after=0,
)
_, tf = textbox(s, Inches(0.8), Inches(6.3), Inches(11.5), Inches(0.4))
para(tf, "Version 2.1  |  August 2026", size=13, color=MINT, space_after=0)

# ===== 2 Agenda =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Agenda")
items = [
    "1.  What is ApnaMedi?",
    "2.  Problems we solve",
    "3.  Platform architecture",
    "4.  Core modules & capabilities",
    "5.  Patient portal",
    "6.  Roles & security",
    "7.  Subscription plans",
    "8.  Live domains & deployment",
    "9.  Business benefits & close",
]
_, tf = textbox(s, Inches(1.2), Inches(1.6), Inches(10), Inches(5))
for i, item in enumerate(items):
    para(tf, item, size=20, color=INK, space_before=8 if i else 0, space_after=4)
footer(s, 2)

# ===== 3 What is =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "What is ApnaMedi?", "One platform for modern healthcare operations")
_, tf = textbox(s, Inches(0.6), Inches(1.5), Inches(12), Inches(1.2))
para(
    tf,
    "ApnaMedi is a cloud-ready multi-tenant Healthcare Management SaaS platform. "
    "One installation serves many independent clinics and hospitals — each with its own "
    "data, branding, staff, and subscription.",
    size=16,
    color=INK,
    space_after=0,
)
cards = [
    ("Multi-tenant SaaS", "Many organizations on one platform with full data isolation"),
    ("Module-based", "Clinic, Laboratory & Diagnostics — enable what you need"),
    ("Role-based access", "Staff see only what their role and plan allow"),
    ("Subscription gated", "Plans control features, limits, and AI add-ons"),
]
for i, (t, d) in enumerate(cards):
    left = Inches(0.5 + i * 3.15)
    add_rect(s, left, Inches(3.0), Inches(3.0), Inches(2.8), WHITE)
    add_rect(s, left, Inches(3.0), Inches(3.0), Inches(0.12), TEAL)
    _, tf = textbox(s, left + Inches(0.2), Inches(3.3), Inches(2.6), Inches(2.3))
    para(tf, t, size=16, bold=True, color=TEAL_DARK, space_after=10)
    para(tf, d, size=13, color=MUTED, space_after=0)
footer(s, 3)

# ===== 4 Problems =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Problems We Solve", "From scattered tools to one connected system")
problems = [
    ("Patient records", "Spreadsheets, paper files, incomplete history"),
    ("Appointments", "Phone booking chaos, no shared calendar"),
    ("Lab & Diagnostics", "Disconnected orders, delays in reports"),
    ("Billing & Finance", "Manual invoices, weak P&L visibility"),
    ("Staff access", "Everyone sees everything — or nothing"),
    ("Multi-branch", "No central control across locations"),
]
for i, (t, d) in enumerate(problems):
    row, col = divmod(i, 3)
    left = Inches(0.5 + col * 4.2)
    top = Inches(1.55 + row * 2.5)
    add_rect(s, left, top, Inches(4.0), Inches(2.2), WHITE)
    add_rect(s, left, top, Inches(0.12), Inches(2.2), TEAL)
    _, tf = textbox(s, left + Inches(0.35), top + Inches(0.35), Inches(3.4), Inches(1.6))
    para(tf, t, size=18, bold=True, color=TEAL_DARK, space_after=12)
    para(tf, d, size=14, color=MUTED, space_after=0)
footer(s, 4)

# ===== 5 Architecture =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Platform Architecture", "Staff app + Patient portal + Shared API")
layers = [
    ("Staff Dashboard", "app.apnamedi.com", "Clinic operations, admin, lab, diagnostics, finance"),
    ("Patient Portal", "patient.apnamedi.com", "Book centres, view reports, prescriptions, chat help"),
    ("Backend API", "app.apnamedi.com/api", "Laravel + Sanctum · Multi-tenant · Role & plan gates"),
]
for i, (t, host, d) in enumerate(layers):
    top = Inches(1.55 + i * 1.6)
    add_rect(s, Inches(0.6), top, Inches(12.1), Inches(1.4), WHITE)
    add_rect(s, Inches(0.6), top, Inches(0.15), Inches(1.4), TEAL)
    _, tf = textbox(s, Inches(1.0), top + Inches(0.25), Inches(11.4), Inches(1.0))
    para(tf, t, size=20, bold=True, color=TEAL_DARK, space_after=4)
    para(tf, f"{host}   ·   {d}", size=14, color=MUTED, space_after=0)
footer(s, 5)

# ===== 6 Modules =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Core Modules & Capabilities")
mods = [
    ("Clinic", ["Patients & chart", "Doctors & schedule", "Appointments", "Prescriptions", "Billing"]),
    ("Laboratory", ["Test catalog", "Packages", "Sample workflow", "Results & verify", "Lab doctors"]),
    (
        "Diagnostics",
        ["Imaging orders", "Packages & discounts", "Referral partners", "Today queue", "Reports & refunds"],
    ),
    ("Platform", ["Multi-branch", "Roles & users", "Finance & P&L", "Audit trail", "Subscriptions"]),
]
for i, (t, items_list) in enumerate(mods):
    left = Inches(0.4 + i * 3.2)
    add_rect(s, left, Inches(1.5), Inches(3.05), Inches(5.2), WHITE)
    add_rect(s, left, Inches(1.5), Inches(3.05), Inches(0.7), TEAL)
    _, tf = textbox(s, left + Inches(0.15), Inches(1.6), Inches(2.75), Inches(0.5))
    para(tf, t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, space_after=0)
    _, tf = textbox(s, left + Inches(0.25), Inches(2.4), Inches(2.6), Inches(4))
    for j, it in enumerate(items_list):
        para(tf, "•  " + it, size=14, color=INK, space_before=8 if j else 0, space_after=2)
footer(s, 6)

# ===== 7 Patient portal =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Patient Portal", "patient.apnamedi.com — self-service for patients")
feats = [
    ("Find centres", "Browse and select diagnostic / clinic centres"),
    ("Book visits", "Choose service, slot, and confirm booking"),
    ("My bookings", "View, cancel, or reschedule appointments"),
    ("Reports", "Access shared lab & diagnostic reports"),
    ("Prescriptions", "View e-prescriptions and medicine details"),
    ("Chat help", "Floating assistant for common patient questions"),
]
for i, (t, d) in enumerate(feats):
    row, col = divmod(i, 3)
    left = Inches(0.5 + col * 4.2)
    top = Inches(1.55 + row * 2.5)
    add_rect(s, left, top, Inches(4.0), Inches(2.2), WHITE)
    add_rect(s, left, top, Inches(4.0), Inches(0.1), TEAL)
    _, tf = textbox(s, left + Inches(0.3), top + Inches(0.4), Inches(3.4), Inches(1.5))
    para(tf, t, size=18, bold=True, color=TEAL_DARK, space_after=10)
    para(tf, d, size=14, color=MUTED, space_after=0)
footer(s, 7)

# ===== 8 Roles =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Roles & Security", "Right access for every team member")
_, tf = textbox(s, Inches(0.6), Inches(1.45), Inches(12), Inches(0.5))
para(
    tf,
    "Effective Access  =  Role Permissions  ∩  Company Modules  ∩  Subscription Features",
    size=15,
    bold=True,
    color=TEAL_DARK,
    space_after=0,
)
roles = [
    ("Super Admin", "Platform owner — companies, plans, theme"),
    ("Company Admin", "Org admin — staff, settings, all modules"),
    ("Doctor", "Own appointments, patients, schedules"),
    ("Receptionist", "Front desk — patients & bookings"),
    ("Lab Technician", "Lab catalog & order workflow"),
    ("Radiologist", "Diagnostic orders & reports"),
    ("Accountant", "Billing & finance access"),
    ("Pharmacist / Nurse", "Clinical & medicine support"),
]
for i, (t, d) in enumerate(roles):
    row, col = divmod(i, 2)
    left = Inches(0.5 + col * 6.4)
    top = Inches(2.15 + row * 1.1)
    add_rect(s, left, top, Inches(6.15), Inches(0.95), WHITE)
    _, tf = textbox(s, left + Inches(0.25), top + Inches(0.15), Inches(5.7), Inches(0.7))
    para(tf, t, size=15, bold=True, color=TEAL_DARK, space_after=2)
    para(tf, d, size=12, color=MUTED, space_after=0)
footer(s, 8)

# ===== 9 Plans =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Subscription Plans", "Grow from small clinic to AI-powered hospital")
plans = [
    ("Basic", "Small clinics &\ndiagnostic centres", "Core ops\nEssential limits"),
    ("Premium", "Growing clinics\n& laboratories", "More modules\nHigher limits"),
    ("Enterprise", "Hospitals &\nmulti-branch", "Full scale\nMulti-branch"),
    ("AI Gold", "AI-powered\nhealthcare", "OCR · Chat\nVoice · Insights"),
]
for i, (t, aud, feat) in enumerate(plans):
    left = Inches(0.45 + i * 3.2)
    add_rect(s, left, Inches(1.55), Inches(3.05), Inches(5.1), WHITE)
    hdr = TEAL_DARK if i < 3 else NAVY
    add_rect(s, left, Inches(1.55), Inches(3.05), Inches(1.1), hdr)
    _, tf = textbox(s, left + Inches(0.15), Inches(1.75), Inches(2.75), Inches(0.7))
    para(tf, t, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER, space_after=0)
    _, tf = textbox(s, left + Inches(0.25), Inches(2.9), Inches(2.6), Inches(3.4))
    para(tf, aud, size=14, color=MUTED, align=PP_ALIGN.CENTER, space_after=16)
    para(tf, feat, size=15, bold=True, color=INK, align=PP_ALIGN.CENTER, space_after=0)
footer(s, 9)

# ===== 10 Live domains =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Live Domains & Deployment", "How production is structured")

add_rect(s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(5.0), WHITE)
add_rect(s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(0.6), TEAL)
_, tf = textbox(s, Inches(0.7), Inches(1.65), Inches(5.6), Inches(0.45))
para(tf, "Production URLs", size=18, bold=True, color=WHITE, space_after=0)
_, tf = textbox(s, Inches(0.8), Inches(2.4), Inches(5.4), Inches(3.8))
para(tf, "Staff app", size=14, bold=True, color=TEAL_DARK, space_after=2)
para(tf, "https://app.apnamedi.com", size=16, color=INK, space_after=14)
para(tf, "Patient portal", size=14, bold=True, color=TEAL_DARK, space_after=2)
para(tf, "https://patient.apnamedi.com", size=16, color=INK, space_after=14)
para(tf, "API", size=14, bold=True, color=TEAL_DARK, space_after=2)
para(tf, "https://app.apnamedi.com/api", size=16, color=INK, space_after=14)
para(tf, "Same frontend build · hostname selects UI", size=13, color=MUTED, space_after=0)

add_rect(s, Inches(6.9), Inches(1.55), Inches(5.9), Inches(5.0), WHITE)
add_rect(s, Inches(6.9), Inches(1.55), Inches(5.9), Inches(0.6), TEAL_DARK)
_, tf = textbox(s, Inches(7.1), Inches(1.65), Inches(5.5), Inches(0.45))
para(tf, "Go-live checklist", size=18, bold=True, color=WHITE, space_after=0)
_, tf = textbox(s, Inches(7.2), Inches(2.4), Inches(5.4), Inches(3.8))
checks = [
    "DNS: patient CNAME/A -> same as app",
    "SSL for both subdomains (Cloudflare)",
    "Serve SPA dist on both hosts",
    "VITE_API_BASE_URL = app.../api",
    "No VITE_APP_MODE on live build",
    "Laravel APP_URL + Sanctum domains",
    "php artisan config:clear",
]
for j, c in enumerate(checks):
    para(tf, "✓  " + c, size=14, color=INK, space_before=6 if j else 0, space_after=2)
footer(s, 10)

# ===== 11 Benefits =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), LIGHT_BG)
section_header(s, "Business Benefits")
bens = [
    ("Faster operations", "Digitize patients, bookings, lab & diagnostics in one place"),
    ("Lower admin load", "Less paperwork, fewer tools, clearer staff workflows"),
    ("Better control", "Roles, audit trail, and subscription limits"),
    ("Patient convenience", "Self-booking, reports, prescriptions online"),
    ("Scalable SaaS", "Onboard many clinics without rebuilding"),
    ("Ready for AI", "AI Gold plan path for OCR, chat & voice"),
]
for i, (t, d) in enumerate(bens):
    row, col = divmod(i, 3)
    left = Inches(0.5 + col * 4.2)
    top = Inches(1.55 + row * 2.5)
    add_rect(s, left, top, Inches(4.0), Inches(2.2), WHITE)
    _, tf = textbox(s, left + Inches(0.3), top + Inches(0.4), Inches(3.4), Inches(1.5))
    para(tf, t, size=18, bold=True, color=TEAL_DARK, space_after=10)
    para(tf, d, size=14, color=MUTED, space_after=0)
footer(s, 11)

# ===== 12 Thank you =====
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), TEAL_DARK)
add_rect(s, Inches(0), Inches(0), Inches(0.25), Inches(7.5), TEAL)
_, tf = textbox(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(1))
para(tf, "Thank You", size=48, bold=True, color=WHITE, space_after=0)
_, tf = textbox(s, Inches(0.8), Inches(3.4), Inches(11.5), Inches(0.6))
para(tf, "ApnaMedi — Healthcare Operations, Simplified", size=22, color=MINT, space_after=0)
_, tf = textbox(s, Inches(0.8), Inches(4.4), Inches(11.5), Inches(1.2))
para(tf, "app.apnamedi.com", size=16, color=WHITE, space_after=6)
para(tf, "patient.apnamedi.com", size=16, color=WHITE, space_after=6)
para(tf, "Questions welcome", size=14, color=MINT_SOFT, space_after=0)

out = r"d:\xamp\htdocs\medeasy\docs\ApnaMedi-Product-Overview.pptx"
prs.save(out)
print(out)
